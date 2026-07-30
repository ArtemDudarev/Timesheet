import io

from openpyxl import Workbook
from openpyxl.styles import Font

from src.schemas.report import PlanFactRow, ProjectStatRow, UtilizationRow

EXPORT_TYPES = ("utilization", "projects", "plan_fact")

_HEADERS = {
    "utilization": ["Сотрудник", "Отработано, ч", "Норма, ч", "Загрузка, %"],
    "projects": ["Проект", "Клиент", "Факт, ч", "Бюджет, ч", "Прогресс, %", "Статус", "Дедлайн"],
    "plan_fact": ["Отдел", "План, ч", "Факт, ч"],
}

_TITLES = {
    "utilization": "Загрузка сотрудников",
    "projects": "Проекты",
    "plan_fact": "План/факт по отделам",
}


def build_xlsx(
    export_type: str,
    period: str,
    rows: list[UtilizationRow] | list[ProjectStatRow] | list[PlanFactRow],
) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = period

    ws.append([f"{_TITLES[export_type]} — {period}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    ws.append(_HEADERS[export_type])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for row in rows:
        if export_type == "utilization":
            ws.append([row.full_name, row.logged_hours, row.norm_hours, row.utilization_pct])
        elif export_type == "projects":
            ws.append([
                row.name, row.client, row.spent_hours, row.budget_hours,
                row.progress_pct, row.status,
                row.deadline.isoformat() if row.deadline else None,
            ])
        else:
            ws.append([row.department_name, row.plan_hours, row.fact_hours])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
